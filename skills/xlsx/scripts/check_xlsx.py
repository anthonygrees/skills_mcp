#!/usr/bin/env python3
"""Summarize an .xlsx/.xlsm workbook and check it for common problems.

Usage:
    python check_xlsx.py book.xlsx            # summary + findings; exit 1 on errors/warnings
    python check_xlsx.py book.xlsx --json     # machine-readable output

What it checks:
  - error values in cells (#REF!, #DIV/0!, #VALUE!, #N/A, #NAME?, #NUM!, #NULL!)
  - formulas that have no cached value (workbook was never recalculated)
  - formulas that point at sheets that do not exist
  - numbers stored as text
  - hardcoded numbers inside formulas (reported as notes, not failures)

Cached values only exist once Excel or LibreOffice has calculated and saved the
file. A workbook written by openpyxl has none, so "never recalculated" is
expected until you recalculate it.

Requires: pip install openpyxl
"""
import argparse
import json
import re
import sys

from openpyxl import load_workbook

ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}
SHEET_REF = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_.]+))!")
NUMERIC_TEXT = re.compile(r"^\s*[-+]?\(?\$?\s?\d[\d,]*(\.\d+)?\)?%?\s*$")
STRING_LITERAL = re.compile(r'"[^"]*"')
CELL_REF = re.compile(r"\$?[A-Z]{1,3}\$?\d+")
LITERAL = re.compile(r"(?<![A-Za-z0-9_.$!:])(\d+\.?\d*)(?![A-Za-z0-9_(!:])")
BENIGN = {"0", "1", "2", "-1", "12", "100", "1000"}


def scan(path):
    keep_vba = path.lower().endswith(".xlsm")
    wf = load_workbook(path, keep_vba=keep_vba)                   # formulas
    wv = load_workbook(path, data_only=True, keep_vba=keep_vba)   # cached values
    sheets = set(wf.sheetnames)
    out = {"file": path, "sheets": [], "errors": [], "warnings": [], "notes": []}
    total_formulas = uncomputed = 0

    for ws in wf.worksheets:
        vs = wv[ws.title]
        info = {"name": ws.title, "dimensions": ws.dimensions, "state": ws.sheet_state,
                "formulas": 0, "merged": len(ws.merged_cells.ranges), "charts": len(ws._charts),
                "images": len(ws._images)}
        text_numbers = []
        notes_for_sheet = 0
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                loc = f"{ws.title}!{c.coordinate}"
                if c.data_type == "f" or (isinstance(v, str) and v.startswith("=")):
                    info["formulas"] += 1
                    total_formulas += 1
                    formula = v if isinstance(v, str) else getattr(v, "text", str(v))
                    cached = vs[c.coordinate].value
                    if cached is None:
                        uncomputed += 1
                    elif isinstance(cached, str) and cached in ERRORS:
                        out["errors"].append(f"{loc}: {cached} from {formula}")
                    for m in SHEET_REF.finditer(formula):
                        name = (m.group(1) or m.group(2) or "").replace("''", "'")
                        if name and name not in sheets and not name.startswith("["):
                            out["errors"].append(f"{loc}: refers to missing sheet '{name}'")
                    bare = CELL_REF.sub("", STRING_LITERAL.sub("", formula))
                    lits = [x for x in LITERAL.findall(bare) if x not in BENIGN]
                    if lits and notes_for_sheet < 5:
                        out["notes"].append(f"{loc}: hardcoded number(s) {sorted(set(lits))} in {formula[:60]}")
                        notes_for_sheet += 1
                elif c.data_type == "e" or (isinstance(v, str) and v in ERRORS):
                    out["errors"].append(f"{loc}: error value {v}")
                elif isinstance(v, str) and NUMERIC_TEXT.match(v):
                    text_numbers.append(c.coordinate)
        if text_numbers:
            shown = ", ".join(text_numbers[:5]) + (" ..." if len(text_numbers) > 5 else "")
            out["warnings"].append(f"{ws.title}: {len(text_numbers)} numbers stored as text ({shown})")
        out["sheets"].append(info)

    if total_formulas and uncomputed == total_formulas:
        out["warnings"].append(
            f"{total_formulas} formulas have no cached values: the workbook has never been recalculated, "
            "so results cannot be checked for errors yet")
    elif uncomputed:
        out["warnings"].append(f"{uncomputed} of {total_formulas} formulas have no cached value")
    out["total_formulas"] = total_formulas
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    res = scan(a.path)
    bad = bool(res["errors"] or res["warnings"])
    if a.json:
        print(json.dumps(res, indent=2))
    else:
        print(f"{res['file']}: {len(res['sheets'])} sheet(s), {res['total_formulas']} formula(s)")
        for s in res["sheets"]:
            extra = []
            if s["state"] != "visible": extra.append(s["state"])
            if s["merged"]: extra.append(f"{s['merged']} merged")
            if s["charts"]: extra.append(f"{s['charts']} chart(s)")
            if s["images"]: extra.append(f"{s['images']} image(s)")
            print(f"  {s['name']}: {s['dimensions']}, {s['formulas']} formulas" + (f" [{', '.join(extra)}]" if extra else ""))
        for label, key in (("ERRORS", "errors"), ("WARNINGS", "warnings"), ("NOTES", "notes")):
            if res[key]:
                print(f"\n{label}")
                for x in res[key][:50]:
                    print(f"  - {x}")
                if len(res[key]) > 50:
                    print(f"  ... {len(res[key]) - 50} more")
        print(f"\n{len(res['errors'])} error(s), {len(res['warnings'])} warning(s), {len(res['notes'])} note(s)")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
